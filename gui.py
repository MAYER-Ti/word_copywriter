from PyQt5 import QtWidgets
from PyQt5.QtWidgets import QFileDialog, QMessageBox
from PyQt5.QtCore import QSettings
from io import BytesIO
import os

from openpyxl import load_workbook

from parsers import read_data_from_file
from doc_utils import format_preview, replace_placeholders


class MainWindow(QtWidgets.QMainWindow):
    def __init__(self):
        super().__init__()
        self.settings = QSettings("word_copywriter", "templates")
        self.templates = {}
        self.data = {}
        self.status_label = QtWidgets.QLabel()
        self.init_ui()
        self.load_templates()

    def init_ui(self):
        self.setWindowTitle("Word Copywriter")

        # Toolbar with settings and status
        toolbar = self.addToolBar("Main")
        settings_button = QtWidgets.QToolButton()
        settings_button.setText("Настройки")
        settings_menu = QtWidgets.QMenu()
        act_action = settings_menu.addAction("Загрузить шаблон акта")
        act_action.triggered.connect(self.browse_act_template)
        invoice_action = settings_menu.addAction("Загрузить шаблон счёта")
        invoice_action.triggered.connect(self.browse_invoice_template)
        settings_button.setMenu(settings_menu)
        settings_button.setPopupMode(QtWidgets.QToolButton.InstantPopup)
        toolbar.addWidget(settings_button)
        toolbar.addSeparator()
        toolbar.addWidget(self.status_label)

        # Central widget layout
        central = QtWidgets.QWidget()
        self.setCentralWidget(central)
        layout = QtWidgets.QVBoxLayout()

        source_layout = QtWidgets.QHBoxLayout()
        source_layout.addWidget(QtWidgets.QLabel("Source"))
        self.source_edit = QtWidgets.QLineEdit()
        self.source_edit.setReadOnly(True)
        source_layout.addWidget(self.source_edit)
        source_btn = QtWidgets.QPushButton("Browse")
        source_btn.clicked.connect(self.browse_source)
        source_layout.addWidget(source_btn)
        layout.addLayout(source_layout)

        self.preview_edit = QtWidgets.QTextEdit()
        self.preview_edit.setReadOnly(True)
        layout.addWidget(self.preview_edit)

        buttons_layout = QtWidgets.QHBoxLayout()
        self.create_act_btn = QtWidgets.QPushButton("Создать акт")
        self.create_act_btn.clicked.connect(self.create_act)
        self.create_act_btn.setEnabled(False)
        buttons_layout.addWidget(self.create_act_btn)

        self.create_invoice_btn = QtWidgets.QPushButton("Создать счёт")
        self.create_invoice_btn.clicked.connect(self.create_invoice)
        self.create_invoice_btn.setEnabled(False)
        buttons_layout.addWidget(self.create_invoice_btn)
        layout.addLayout(buttons_layout)
        
        central.setLayout(layout)

    def set_status(self, message):
        self.status_label.setText(message)

    def load_template(self, key):
        path = self.settings.value(f"{key}_template", "", type=str)
        name = "акта" if key == "act" else "счёта"
        if path and os.path.exists(path):
            try:
                with open(path, "rb") as f:
                    self.templates[key] = f.read()
                self.set_status(f"Шаблон {name} загружен")
            except Exception as e:
                self.templates[key] = None
                self.set_status(f"Ошибка загрузки {name}: {e}")
        else:
            self.templates[key] = None

    def load_templates(self):
        self.templates = {}
        for key in ("act", "invoice"):
            self.load_template(key)
        self.update_create_buttons_state()

    def browse_source(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Select source", filter="Documents (*.docx *.pdf)"
        )
        if path:
            self.source_edit.setText(path)
            self.data = read_data_from_file(path)
            self.preview_edit.setPlainText(format_preview(self.data))
        self.update_create_buttons_state()

    def browse_act_template(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Select act template", filter="Excel Files (*.xls)"
        )
        if path:
            self.settings.setValue("act_template", path)
            self.load_template("act")
        self.update_create_buttons_state()

    def browse_invoice_template(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Select invoice template", filter="Excel Files (*.xls)"
        )
        if path:
            self.settings.setValue("invoice_template", path)
            self.load_template("invoice")
        self.update_create_buttons_state()

    def create_act(self):
        self.create_document("act")

    def create_invoice(self):
        self.create_document("invoice")

    def create_document(self, key):
        template_bytes = self.templates.get(key)
        source_path = self.source_edit.text()
        if not (template_bytes and source_path):
            QMessageBox.warning(self, "Warning", "Please select source and templates")
            return
        output_path, _ = QFileDialog.getSaveFileName(
            self, "Save document", filter="Excel Files (*.xlsx)"
        )
        if not output_path:
            return
        if not output_path.lower().endswith(".xlsx"):
            output_path += ".xlsx"
        data = getattr(self, "data", None)
        if not data:
            data = read_data_from_file(source_path)
        try:
            wb = load_workbook(BytesIO(template_bytes))
            replace_placeholders(wb, data)
            wb.save(output_path)
            QMessageBox.information(self, "Success", f"Document saved to {output_path}")
        except PermissionError:
            QMessageBox.critical(
                self,
                "Error",
                "Cannot save file. It may be open in another program.",
            )
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Failed to save file: {e}")

    def update_create_buttons_state(self):
        source_selected = bool(self.source_edit.text())
        self.create_act_btn.setEnabled(source_selected and bool(self.templates.get("act")))
        self.create_invoice_btn.setEnabled(
            source_selected and bool(self.templates.get("invoice"))
        )
